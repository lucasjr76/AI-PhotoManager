"""PDF (pypdfium2) and DOCX (python-docx) text and metadata. Read-only."""

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

import docx
import pypdfium2 as pdfium

from aipdm.core.dates import parse_pdf_date
from aipdm.core.images import THUMB_SIZE, save_thumbnail


@dataclass(frozen=True)
class DocumentInfo:
    pages: list[tuple[int, str]]  # (1-based page, text); pages without a text layer are skipped
    created: datetime | None


def read_pdf(path: Path, thumb_path: Path) -> DocumentInfo:
    # Open from bytes, not path: pdfium never gets a handle on the source file.
    pdf = pdfium.PdfDocument(path.read_bytes())
    try:
        pages: list[tuple[int, str]] = []
        for index in range(len(pdf)):
            page = pdf[index]
            text = page.get_textpage().get_text_bounded().strip()
            if text:
                pages.append((index + 1, text))
            if index == 0:
                scale = THUMB_SIZE / max(page.get_size())
                save_thumbnail(page.render(scale=scale).to_pil(), thumb_path)
        created = parse_pdf_date(str(pdf.get_metadata_dict().get("CreationDate", "")))
        return DocumentInfo(pages, created)
    finally:
        pdf.close()


def read_docx(path: Path) -> DocumentInfo:
    with path.open("rb") as fh:
        document = docx.Document(fh)
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    text = "\n".join(p for p in parts if p.strip())
    created = document.core_properties.created
    if created is not None and created.tzinfo is not None:
        created = created.astimezone().replace(tzinfo=None)
    return DocumentInfo([(1, text)] if text else [], created)
