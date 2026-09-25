"""Small fixtures generated on the fly; no personal files are versioned."""

import os
from collections.abc import Iterator
from functools import partial
from pathlib import Path

import docx
import pytest
from PIL import Image

from aipdm.core.paths import models_dir
from aipdm.core.scanner import index

# Phase-1 stages only (date, thumb, text): no models needed.
index_base = partial(index, only=frozenset())

MODELS_PRESENT = (models_dir() / "clip_image.onnx").exists() and (
    models_dir() / "latin_PP-OCRv5_rec_mobile.onnx"
).exists()
requires_models = pytest.mark.skipif(
    not MODELS_PRESENT, reason="modelos ausentes: rode tools/fetch_models.py e export_clip.py"
)

EXIF_IFD = 0x8769
DATETIME_ORIGINAL = 0x9003


def make_pdf(path: Path, pages: list[str], creation_date: str = "D:20210304050607-03'00'") -> None:
    """Hand-built PDF with a real text layer (Helvetica, WinAnsi so 'ã' survives)."""
    objs: list[bytes] = []
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        content = f"BT /F1 18 Tf 50 700 Td ({text}) Tj ET".encode("cp1252")
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {4 + 2 * i} 0 R"
            f" /Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objs.append(b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream")
    objs.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )
    objs.append(f"<< /CreationDate ({creation_date}) >>".encode())
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % n + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R /Info %d 0 R >>\n" % (len(objs) + 1, len(objs))
    out += b"startxref\n%d\n%%%%EOF\n" % xref
    path.write_bytes(bytes(out))


def make_docx(path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Relatório escrito por João da Silva.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cidade"
    table.rows[0].cells[1].text = "Florianópolis"
    document.save(str(path))


def build_sample(root: Path) -> None:
    images = root / "WhatsApp Images"
    images.mkdir(parents=True)
    Image.new("RGB", (64, 48), "red").save(images / "IMG-20230514-WA0001.jpg")
    Image.new("RGB", (30, 60), "blue").save(images / "WhatsApp Image 2022-01-02 at 9.05.09 PM.png")

    exif = Image.Exif()
    exif.get_ifd(EXIF_IFD)[DATETIME_ORIGINAL] = "2019:07:08 10:11:12"
    Image.new("RGB", (40, 40), "green").save(root / "camera.jpg", exif=exif)

    stickers = root / "WhatsApp Stickers"
    stickers.mkdir()
    Image.new("RGBA", (32, 32)).save(stickers / "STK-20230101-WA0002.webp")

    docs = root / "WhatsApp Documents"
    docs.mkdir()
    make_pdf(docs / "contrato.pdf", ["Contrato de João e Maria", "Praia de Jurere"])
    make_docx(docs / "relatorio.docx")

    (root / "corrompida.jpg").write_bytes(b"not really a jpeg")
    (root / "VID-20230514-WA0003.mp4").write_bytes(b"\x00" * 128)


@pytest.fixture
def sample_dir(tmp_path: Path) -> Path:
    root = tmp_path / "pasta"
    build_sample(root)
    return root


@pytest.fixture(autouse=True)
def isolated_data_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    data = tmp_path / "dados"
    monkeypatch.setenv("AIPDM_DATA_DIR", str(data))
    return data


def set_tree_readonly(root: Path, readonly: bool) -> None:
    for dirpath, _dirs, files in os.walk(root):
        for name in files:
            os.chmod(Path(dirpath) / name, 0o444 if readonly else 0o644)
        os.chmod(dirpath, 0o555 if readonly else 0o755)


@pytest.fixture
def readonly_sample(sample_dir: Path) -> Iterator[Path]:
    set_tree_readonly(sample_dir, True)
    yield sample_dir
    set_tree_readonly(sample_dir, False)  # let pytest clean tmp_path
